from __future__ import annotations

import json
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (BaseDocTemplate, Frame, PageBreak, PageTemplate, Paragraph,
                                Spacer, Table, TableStyle, KeepTogether)

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"
OUT.mkdir(exist_ok=True)

PAGES = [
    ("Executive summary", [
        ("p", "Human API helps handle a work question when the colleague who knows the answer is unavailable. The colleague chooses what questions the software may handle and for how long. The person asking must still be allowed to open every file used. The app shows evidence and what remains unknown, or sends the question to a human."),
        ("p", "Example: Alice asks why a service timeout changed from 90 to 120 seconds. Bob has opted in for Project Alpha questions. The app finds a written technical decision, ADR-024, that Alice can read. It explains why the change was approved for release 5.2, but not whether 120 seconds is still needed today. The app gives that partial answer with its source and limitation; it does not pretend Bob wrote it."),
        ("p", "This is an early hackathon concept and a local working demo. People, projects and documents are fictional. The answer uses word matching and fixed software rules, not a language model (LLM). Microsoft sign-in and file-search adapters are not connected to a real tenant; there is no Teams bot."),
        ("callout", "Product hypothesis: make an unanswered question a consented, time-bounded, evidence-backed and auditable work item. Search and access-aware AI are established capabilities; Human API should integrate with Microsoft Copilot rather than duplicate enterprise search."),
        ("p", "Current answer generation is deterministic evidence extraction, not an LLM. The Graph and Entra adapters are code scaffolding and have not been tested against a tenant; Teams messaging is not implemented. No production readiness, accreditation or measured Presight productivity gain is claimed."),
    ]),
    ("Problem, people and scenario", [
        ("h", "The workplace dependency problem"),
        ("p", "Work stalls when an employee needs a decision rationale, operational detail or project history held by a colleague who is busy, offline, on leave or no longer assigned. The knowledge may already be in an ADR, runbook or approved project source, but finding and validating it still takes active time. Waiting time is not automatically productive time saved."),
        ("h", "Fictional scenario"),
        ("p", "Alice, a platform engineer, asks Bob why Project Alpha's startup timeout changed to 120 seconds and whether the change is still needed. Bob is on fictional leave and has approved a time-limited policy for this project, this topic, selected files and named coworkers. ADR-024 is a written technical decision record. It explains that the change was approved for release 5.2 because dependency setup sometimes took more than 90 seconds. It does not say whether 120 seconds is still needed. Human API shows this partial answer with the source and limitation; it does not speak as Bob."),
        ("h", "Users"),
        ("bullet", "Requester: submit a question, view an answer and evidence, give feedback or request review."),
        ("bullet", "Knowledge owner: opt in, choose dates, projects, sources, topics, audience and backup; revoke at any time."),
        ("bullet", "Manager: approve scope changes and monitor team blocking patterns without unrestricted access to private question content."),
        ("bullet", "Security administrator: manage policy and audit metadata; no default ability to read every question or source."),
    ]),
    ("Solution and workflow", [
        ("h", "Lifecycle"),
        ("bullet", "Save the question, who asked it, who may know, the project and the response deadline."),
        ("bullet", "Wait for the agreed response time, unless an owner-approved out-of-office rule applies. Busy status alone never triggers the helper."),
        ("bullet", "Check the owner's limits: dates, project, topic, approved coworkers, files and manager review."),
        ("bullet", "Before searching a file, check that the person asking can already open it."),
        ("bullet", "Give a source-backed answer and its limits, deny a restricted request without exposing details, or send a question to a human."),
        ("bullet", "Label a human reply as human. Require approval before reusing a correction."),
        ("h", "MVP screens"),
        ("p", "The React app contains a dashboard, ask flow, My AI Delegate controls, escalation inbox and governance view. It uses real local API results. Identity and availability are explicitly synthetic. The interface labels software-produced text as AI-generated to distinguish it from a later human response; the demo uses fixed templates, not a language model."),
        ("callout", "Safety outcome: lack of evidence is a reason to abstain. It is not an invitation to fill gaps with plausible text."),
    ]),
    ("Architecture and technical feasibility", [
        ("p", "Developer view: the MVP is a modular monolith—one app organized into separate code modules. The web API calls rule-based policy, access, search and workflow modules. SQLAlchemy saves people, projects, policies, files, questions, status history, deadlines, answers, citations, escalations, reviews and audit events in a relational database."),
        ("p", "A database-backed timeout scanner polls pending jobs; conditional transitions and unique question/job/answer relationships reduce duplicate outcomes in the single-process demo. Production multi-worker operation still needs database-specific claims/locks, an outbox, observability, recovery tests and idempotent delivery."),
        ("p", "The demo finds files by matching words; it does not understand paraphrases or meaning as a language model would. It checks the requester's project membership and file-access list before scoring results. It cites matching text, identifies known gaps or conflicts, and stops when evidence is not enough. It does not use embeddings, a vector database or a language model. This is a deliberate demo limit, not a recommended production search engine."),
        ("h", "Technical stack"),
        ("p", "Python 3.12, FastAPI 0.162.0, Pydantic 2.14.1, SQLAlchemy 2.1.6, Alembic 1.20.1, PostgreSQL 18.6 for Compose, psycopg 3.3.6, React 19.3.1, TypeScript 7.0.2 and Vite 8.3.5. Python versions are pinned; npm lockfiles are included. No model weights or paid AI APIs are bundled."),
        ("p", "LangGraph was not selected: a small deterministic state machine with human handoff and durable rows is sufficient at this scope. Add orchestration only if a measured multi-tool/checkpoint need appears."),
    ]),
    ("Microsoft 365 integration", [
        ("p", "Recommended path: build a Teams or Copilot Studio extension around this question workflow. Reuse Microsoft work sign-in and file permissions instead of building a new company-wide search product."),
        ("bullet", "Work sign-in: early Microsoft Entra ID token-checking code exists. It needs an app registration and tenant testing before it can sign anyone in."),
        ("bullet", "SharePoint: an early Microsoft Graph adapter is designed to search and fetch files as the person asking, preserving that person's file permissions. It has not been tested end to end in a tenant."),
        ("bullet", "Outlook: a possible adapter reads only the signed-in owner's out-of-office status and dates, then discards the reply text. It does not read inbox messages."),
        ("bullet", "Teams: no bot is implemented. A future Teams app needs registration, an internet-reachable secure endpoint and permission to install the app in the tenant."),
        ("p", "Microsoft Graph uses named permissions to control which files an app may search. Some permission choices need administrator approval. Choose the narrowest permission that works in the target tenant. Never use the owner's broader file access to answer someone else's question. Permission, licensing and tenant rules must be tested before a pilot."),
    ]),
    ("Security, privacy and responsible AI", [
        ("h", "Non-negotiable policy"),
        ("bullet", "If identity or permission is unclear, deny the request. Require verified work sign-in and the owner's explicit approval or another approved organizational policy."),
        ("bullet", "The owner's approval never gives another person new file access. Keep the requester's file permissions and check again before showing saved evidence."),
        ("bullet", "Restricted files need a named list of people who may open them; being on the project alone is not enough."),
        ("bullet", "Do not reveal hidden titles/snippets through denials, summaries, caches or admin views. Keep security administration separate from content access."),
        ("bullet", "Treat file contents as information, not commands. The demo filters a few known hostile phrases; a real service needs broader attack testing and carefully limited tools."),
        ("bullet", "Never impersonate the owner. Mark AI text as generated. Human corrections require publication approval."),
        ("p", "Synthetic labels are examples; they are not Presight's official sensitivity categories. The app does not include encrypted database keys, production data deletion rules, data-loss-prevention labels, incident response or security certification. Use fictional data until a security/privacy review and pilot controls are complete."),
    ]),
    ("Competitive position", [
        ("p", "Microsoft 365 Copilot, Copilot Studio, SharePoint agents, Glean and Moveworks already offer enterprise search or grounded answers, citations, permission-aware sources, configurable agents and/or human handoff. Microsoft announced Autopilot (formerly Scout) on 25 September 2026 as a persistent proactive agent with a tenant identity; the announcement said private preview was expanding at the end of September, not general availability. Microsoft also announced a Copilot/agent-curated team knowledge base in SharePoint Frontier Preview on 14 October, while Copilot in SharePoint began its GA rollout on 30 September. This substantially overlaps broad “knowledge continuity” messaging."),
        ("p", "A plausible differentiating workflow is narrower: an employee-consented and dated delegation; unanswered request with a configured deadline; project/topic/audience/source scope; requester-side live ACL; evidence-backed answer or abstention; backup escalation; approved correction history. Public materials do not prove that this exact lifecycle is unique. Validate it against Microsoft configurability and customer interviews."),
        ("h", "Recommendation"),
        ("p", "Prototype inside Teams or Copilot Studio first. Use SharePoint agents or Microsoft 365 Copilot for general knowledge grounding; use Human API only for explicit handoff, response-window, approval and audit needs. Reassess against Autopilot preview and vendor capabilities before any standalone investment."),
        ("p", "Glean and Moveworks do not publish comparable public list prices in the reviewed technical sources. Microsoft Copilot/Studio/agent usage depends on eligible user plans, capacity or usage-based billing and tenant setup. Procurement requires current vendor quotes and tenant-specific licensing review."),
    ]),
    ("Implemented MVP and evaluation", [
        ("h", "What works today"),
        ("bullet", "20 fictional employees, 4 teams, 5 projects, 55 approved synthetic documents and a demo policy set."),
        ("bullet", "Persistent questions, timeout jobs, owner consent, manager approval, scope checks, evidence citations, denial, abstention, correction and audit metadata."),
            ("bullet", "17 automated tests pass, including restart durability and idempotency checks; see Evaluation Report for coverage."),
        ("bullet", "36 labeled requester-scoped retrieval cases: 32 expected authorized hits and 4 expected no-evidence cases. All 36 pass; this is lexical retrieval/ACL fixture coverage only, not answer accuracy or security assurance."),
        ("p", "Covered tests include prompt-injection fixture handling, denial without restricted title/snippet, policy revocation, forged demo identity in enterprise mode, owner approval, deadline cancellation, retry limits and human-before-timeout race. No live Microsoft API, LLM evaluation, performance benchmark or real-user pilot has been run."),
        ("h", "MVP acceptance boundary"),
        ("p", "The local app runs with documented commands and synthetic data. Entra/Graph calls need tenant credentials and consent; the enterprise question route remains disabled. Teams messaging, production token management, persistent cache invalidation, multi-instance job leasing and calibrated semantic validation remain future work."),
    ]),
    ("Business value and illustrative time estimate", [
        ("p", "The hoped-for benefit is less active time spent searching, asking again and reconstructing decisions. Count only work time a pilot can show was actually avoided after a useful answer. Do not count hours spent waiting as saved work by default."),
        ("p", "Example math, not a Presight measurement: 500 people × 2 blocking questions each week × 50 weeks = 50,000 possible questions a year. The scenarios below then estimate what share is suitable, used and resolved, and how many active minutes are saved. The AED 250 hourly labor rate is also an assumption."),
        ("table", [
            ["Scenario", "Suitable × used × solved", "Active min saved", "Time value / year"],
            ["Low", "30% × 40% × 50% = 3,000", "10", "AED 125,000"],
            ["Base", "40% × 60% × 70% = 8,400", "20", "AED 700,000"],
            ["High", "55% × 75% × 75% ≈ 15,469", "25", "AED 1,611,328"],
        ]),
        ("p", "Illustrative annual operating cost: AED 740,000 before added Microsoft licenses. This assumes 0.7 of a full-time engineer (AED 336k), 0.4 of a full-time operations/support role (AED 192k), infrastructure and monitoring (AED 60k), security/privacy review (AED 120k), and adoption/support (AED 32k). These are planning placeholders, not quotes. The base estimate of time value is AED 40k below the estimated operating cost, so a pilot must prove value or the cost must fall."),
        ("p", "Potential future pricing could be per user/team, private deployment, custom connectors and support. Price and revenue are hypothetical; no willingness-to-pay evidence exists. M365/agent licensing cost is excluded because it depends on the tenant contract."),
    ]),
    ("Risks, roadmap and references", [
        ("h", "Risks and mitigations"),
        ("bullet", "Showing a file to someone who lost access: check requester permissions before search and again before displaying saved evidence; test revoked access."),
        ("bullet", "Incorrect or stale answer: cite source version/date, show uncertainty, abstain on conflict, require expert approval for reuse."),
        ("bullet", "Employee surveillance concern: explicit opt-in, no inbox reading, no Busy-as-consent, date-bound revocation."),
        ("bullet", "Weak business case or competing Microsoft features: run a short team pilot, compare with existing Copilot setup, and stop if users do not save active time or trust the workflow."),
        ("h", "Next steps"),
        ("p", "0–2 weeks: security/privacy and product discovery; choose one low-sensitivity team. 2–6 weeks: Entra + delegated Graph proof and Teams workflow, with source permissions revalidated. 6–10 weeks: pilot against labeled questions and human acceptance, tracking active work minutes, coverage, correction, escalation and zero access violations. Advance only with approval from Security, Privacy, Microsoft platform owners and pilot users."),
        ("h", "Selected references (accessed 26 October 2026)"),
        ("p", "Microsoft Autopilot announcement, 25 Sep 2026: blogs.microsoft.com/blog/2026/09/25/introducing-the-new-copilot-with-home-code-and-autopilot. Graph Search API: learn.microsoft.com/en-us/graph/api/resources/search-api-overview. Teams SDK registration: learn.microsoft.com/en-us/microsoftteams/platform/teams-sdk/get-started/quickstart-register."),
        ("p", "Microsoft SharePoint agent access: learn.microsoft.com/en-us/sharepoint/manage-access-agents-in-sharepoint. Copilot Studio licensing: learn.microsoft.com/en-us/microsoft-copilot-studio/requirements-licensing. Glean: docs.glean.com/user-guide/about/what-is-glean. Moveworks: help.moveworks.com/ai-assistant/capabilities. PostgreSQL version support: postgresql.org/support/versioning."),
    ]),
]

NAVY = "17352F"
GREEN = "2F8069"
PALE = "EAF2ED"
INK = "24342F"
MUTED = "61746D"
CORAL = "D86E5A"


def docx_base() -> Document:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    for name, size, color in [("Heading 1", 20, NAVY), ("Heading 2", 12, GREEN)]:
        style = document.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(6)
        style.paragraph_format.space_after = Pt(4)
    header = section.header.paragraphs[0]
    header.text = "HUMAN API  /  EUREKA INNOVATION HACKATHON"
    header.style = document.styles["Caption"]
    header.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Concept formulated · 26 October 2026     |     ")
    field = footer.add_run()
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    field._r.addnext(fld)
    return document


def build_docx() -> None:
    document = docx_base()
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(60)
    run = p.add_run("HUMAN API")
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(42)
    run.font.color.rgb = RGBColor.from_string(GREEN)
    subtitle = document.add_paragraph("A governed workflow for work questions when an expert is unavailable")
    subtitle.runs[0].font.size = Pt(20)
    subtitle.runs[0].font.color.rgb = RGBColor.from_string(NAVY)
    document.add_paragraph("Eureka Innovation Hackathon  ·  Future of Work @ Presight\nEarly concept · local demo · fictional data")
    document.add_paragraph("Submission concept and local synthetic-data MVP\n26 October 2026 · Dubai")
    document.add_paragraph("\nA clear, safe way to handle work questions when the expert is unavailable.")
    document.add_heading("What is Human API?", level=2)
    document.add_paragraph("A colleague chooses what questions the software may help with and for how long. The person asking must still be allowed to open every file used. The app shows what a source says and what it does not say, or routes the question to a human.")
    document.add_paragraph("Example: Alice asks why a timeout changed from 90 to 120 seconds. Bob opted in for this project. The app finds a written decision Alice can read. It explains why the change was approved for release 5.2, but cannot tell whether the setting is still needed today. The answer includes the source and that limitation; it is not presented as Bob's reply.")
    document.add_paragraph("This is an early local demo with fictional people and files. It finds evidence by matching words and uses fixed response templates; it does not use a language model. Microsoft sign-in and file-search adapters are untested scaffolding, and no Teams bot exists. See the accompanying Reader Guide for a glossary and AI-agent instructions.")
    document.add_page_break()
    for title, blocks in PAGES[1:]:
        document.add_heading(title, level=1)
        for kind, text in blocks:
            if kind == "h":
                document.add_heading(text, level=2)
            elif kind == "bullet":
                p = document.add_paragraph(style="List Bullet")
                p.add_run(text)
            elif kind == "callout":
                p = document.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.18)
                p.paragraph_format.right_indent = Inches(0.12)
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(8)
                r = p.add_run(text)
                r.bold = True
                r.font.color.rgb = RGBColor.from_string(GREEN)
            elif kind == "table":
                table = document.add_table(rows=0, cols=4)
                table.style = "Light Shading Accent 1"
                for row in text:
                    cells = table.add_row().cells
                    for index, value in enumerate(row):
                        cells[index].text = value
            else:
                document.add_paragraph(text)
        if title != PAGES[-1][0]:
            document.add_page_break()
    document.save(OUT / "HumanAPI_Idea_Document.docx")


def pdf_styles():
    base = getSampleStyleSheet()
    return {
        "body": ParagraphStyle("BodyHA", parent=base["BodyText"], fontName="Helvetica",
                               fontSize=9.2, leading=13.3, textColor=colors.HexColor("#" + INK),
                               spaceAfter=6),
        "bullet": ParagraphStyle("BulletHA", parent=base["BodyText"], fontName="Helvetica",
                                 fontSize=9, leading=12, leftIndent=12, firstLineIndent=-8,
                                 textColor=colors.HexColor("#" + INK), spaceAfter=3),
        "h1": ParagraphStyle("H1HA", parent=base["Heading1"], fontName="Helvetica-Bold",
                            fontSize=18, leading=22, textColor=colors.HexColor("#" + NAVY),
                            spaceBefore=2, spaceAfter=10),
        "h2": ParagraphStyle("H2HA", parent=base["Heading2"], fontName="Helvetica-Bold",
                            fontSize=11, leading=14, textColor=colors.HexColor("#" + GREEN),
                            spaceBefore=6, spaceAfter=4),
        "call": ParagraphStyle("CallHA", parent=base["BodyText"], fontName="Helvetica-Bold",
                              fontSize=10, leading=14, textColor=colors.HexColor("#" + NAVY),
                              backColor=colors.HexColor("#" + PALE), borderPadding=8,
                              spaceBefore=8, spaceAfter=8),
        "small": ParagraphStyle("SmallHA", parent=base["BodyText"], fontName="Helvetica",
                                fontSize=7.2, leading=9.2, textColor=colors.HexColor("#" + MUTED),
                                spaceAfter=3),
        "title": ParagraphStyle("TitleHA", parent=base["Title"], fontName="Helvetica-Bold",
                               fontSize=36, leading=40, textColor=colors.HexColor("#" + NAVY),
                               alignment=TA_CENTER, spaceAfter=16),
    }


def decorate(canvas, doc):
    canvas.saveState()
    width, height = A4
    canvas.setFillColor(colors.HexColor("#" + NAVY))
    canvas.rect(0, height - 11 * mm, width, 11 * mm, fill=1, stroke=0)
    canvas.setFont("Helvetica-Bold", 8)
    canvas.setFillColor(colors.HexColor("#" + PALE))
    canvas.drawString(18 * mm, height - 7.3 * mm, "HUMAN API  /  EUREKA HACKATHON")
    canvas.setStrokeColor(colors.HexColor("#D7E3DC"))
    canvas.line(18 * mm, 15 * mm, width - 18 * mm, 15 * mm)
    canvas.setFont("Helvetica", 7.5)
    canvas.setFillColor(colors.HexColor("#" + MUTED))
    canvas.drawString(18 * mm, 10 * mm, "Concept formulated · 26 October 2026 · synthetic data only")
    canvas.drawRightString(width - 18 * mm, 10 * mm, str(doc.page))
    canvas.restoreState()


def as_flow(block, styles):
    kind, value = block
    if kind == "h":
        return Paragraph(value, styles["h2"])
    if kind == "bullet":
        return Paragraph("• " + value, styles["bullet"])
    if kind == "callout":
        return Paragraph(value, styles["call"])
    if kind == "table":
        return make_table(value, [49 * mm, 54 * mm, 34 * mm, 40 * mm])
    if kind == "p":
        return Paragraph(value.replace("&", "&amp;"), styles["body"])
    return Spacer(1, 2)


def make_table(rows, widths):
    data = [[Paragraph(str(x).replace("&", "&amp;"), ParagraphStyle(
        "cell", fontName="Helvetica-Bold" if ri == 0 else "Helvetica", fontSize=7.6,
        leading=9.5, textColor=colors.white if ri == 0 else colors.HexColor("#" + INK)))
             for x in row] for ri, row in enumerate(rows)]
    table = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + NAVY)),
        ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor("#F2F6F3")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D6E0DA")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    return table


def build_idea_pdf() -> None:
    styles = pdf_styles()
    story = [Spacer(1, 36 * mm), Paragraph("HUMAN API", styles["title"]),
             Paragraph("A governed workflow for work questions when an expert is unavailable",
                       ParagraphStyle("Sub", parent=styles["body"], fontSize=16, leading=22,
                                      alignment=TA_CENTER, textColor=colors.HexColor("#" + GREEN))),
             Spacer(1, 10 * mm), Paragraph("Eureka Innovation Hackathon · Future of Work @ Presight",
                       ParagraphStyle("CenterHA", parent=styles["body"], alignment=TA_CENTER, fontSize=11)),
             Paragraph("Early concept · local demo · fictional data",
                       ParagraphStyle("CenterHA2", parent=styles["body"], alignment=TA_CENTER)),
             Spacer(1, 20 * mm), Paragraph("A clear, safe way to handle work questions when the expert is unavailable.",
                       ParagraphStyle("Lead", parent=styles["call"], alignment=TA_CENTER)),
             Spacer(1, 6 * mm), Paragraph("<b>What is Human API?</b><br/>A colleague chooses what questions the software may help with and for how long. The person asking must still be allowed to open every file used. The app shows what a source says and what it does not say, or routes the question to a human.<br/><br/><b>Example:</b> Alice asks why a timeout changed from 90 to 120 seconds. Bob opted in for this project. The app finds a written decision Alice can read. It explains why the change was approved for release 5.2, but cannot tell whether the setting is still needed today. The answer includes the source and that limitation; it is not presented as Bob's reply.<br/><br/>This is an early local demo with fictional people and files. It finds evidence by matching words and uses fixed response templates; it does not use a language model. Microsoft sign-in and file-search adapters are untested scaffolding, and no Teams bot exists. See the accompanying Reader Guide for a glossary and AI-agent instructions.",
                       ParagraphStyle("SummaryCover", parent=styles["body"], fontSize=10.2, leading=14,
                                      alignment=1, textColor=colors.HexColor("#" + INK))),
             Spacer(1, 10 * mm), Paragraph("Submission concept and synthetic-data MVP · 26 October 2026 · Dubai",
                       ParagraphStyle("Date", parent=styles["small"], alignment=TA_CENTER)), PageBreak()]
    for title, blocks in PAGES[1:]:
        story.append(Paragraph(title, styles["h1"]))
        for block in blocks:
            item = as_flow(block, styles)
            story.append(item)
            story.append(Spacer(1, 3))
        if title != PAGES[-1][0]:
            story.append(PageBreak())
    doc = BaseDocTemplate(str(OUT / "HumanAPI_Idea_Document.pdf"), pagesize=A4,
                          leftMargin=18 * mm, rightMargin=18 * mm,
                          topMargin=18 * mm, bottomMargin=20 * mm)
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="normal")
    doc.addPageTemplates([PageTemplate(id="all", frames=frame, onPage=decorate)])
    doc.build(story)


def simple_pdf(filename: str, title: str, pages: list[tuple[str, list[tuple[str, str]]]]) -> None:
    styles = pdf_styles()
    story = [Spacer(1, 4 * mm), Paragraph(title, styles["h1"]),
             Paragraph("Human API · 26 October 2026 · synthetic concept analysis", styles["small"]),
             Spacer(1, 4 * mm)]
    for index, (heading, blocks) in enumerate(pages):
        if index:
            story.append(PageBreak())
        story.append(Paragraph(heading, styles["h1"]))
        for block in blocks:
            kind, value = block
            if kind == "table":
                story.extend([make_table(value, [31 * mm, 54 * mm, 40 * mm, 48 * mm]), Spacer(1, 4)])
            else:
                story.append(as_flow(block, styles))
                story.append(Spacer(1, 3))
    doc = BaseDocTemplate(str(OUT / filename), pagesize=A4,
                          leftMargin=18 * mm, rightMargin=18 * mm,
                          topMargin=18 * mm, bottomMargin=20 * mm)
    doc.addPageTemplates([PageTemplate(id="all", frames=Frame(
        doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="normal"), onPage=decorate)])
    doc.build(story)


def build_business_case():
    pages = [
        ("1 · Opportunity and hypothesis", [
            ("p", "When people cannot find a decision or reach the colleague who knows it, they search again, send follow-ups or interrupt another expert. Human API would make each question a trackable request and use only files the person asking can already open."),
            ("callout", "The value estimate is time that might become available for other work. It is not cash saved unless the organization actually avoids a cost or produces more with that time."),
            ("h", "Value pathways"),
            ("bullet", "Less active time spent searching and rebuilding context."),
            ("bullet", "Fewer repeated interruptions and clearer handovers."),
            ("bullet", "Easier access to approved technical decisions for newer teammates."),
            ("bullet", "A visible record of who approved the workflow, what source was used and whether a person corrected it."),
            ("p", "This is an unvalidated concept. There is no employee time study, customer feedback or measured company benefit."),
        ]),
        ("2 · Illustrative value of time made available", [
            ("p", "Example assumptions: 500 employees; each has 2 work-blocking questions per week; 50 work weeks a year; and employee time costs AED 250 an hour after salary and employer costs. This creates 50,000 possible questions. None of these figures is a measured Presight baseline."),
            ("p", "For each scenario: ‘suitable’ is the share allowed by the policy and file permissions; ‘used’ is the share employees actually submit; ‘solved’ is the share that gets an answer the requester accepts. Active minutes are time actually spent searching, not time waiting."),
            ("table", [["Case", "Suitable × used × solved", "Minutes saved per answer", "Estimated time value / year"],
                       ["Low", "30% × 40% × 50% (3,000)", "10", "AED 125,000"],
                       ["Base", "40% × 60% × 70% (8,400)", "20", "AED 700,000"],
                       ["High", "55% × 75% × 75% (15,469)", "25", "AED 1,611,328"]]),
            ("p", "Worked example: in the base case, 8,400 questions are accepted. If each avoids 20 minutes of active searching, that is 2,800 hours. At AED 250 per hour, those hours have an estimated value of AED 700,000."),
            ("p", "Formula: 50,000 possible questions × suitable share × used share × solved share × minutes saved ÷ 60 × AED 250/hour. The high case is only a sensitivity example, not a forecast. Do not count idle waiting time or count the same saved interruption twice."),
        ]),
        ("3 · Cost, sensitivity and commercial model", [
            ("table", [["Annual cost assumption", "Basis", "AED"],
                       ["Engineering", "0.7 FTE × AED 480k loaded", "336,000"],
                       ["Operations and support", "0.4 FTE × AED 480k loaded", "192,000"],
                       ["Infrastructure / observability", "Planning estimate", "60,000"],
                       ["Security and privacy review", "Planning estimate", "120,000"],
                       ["Adoption and support", "Planning estimate", "32,000"],
                       ["Total before Microsoft licenses", "Illustrative", "740,000"]]),
            ("p", "The base estimate of AED 700,000 in time value is AED 40,000 below the example AED 740,000 yearly operating cost—and does not include added Microsoft licenses. This is not a break-even forecast: time value may not become cash savings. The demo has no AI model cost. If a language model is added later, include its hosting and usage cost."),
            ("p", "Possible future ways to charge could include a fee per user or team, private hosting, connector setup or support. These are ideas only; no customer willingness to pay or market price was verified."),
        ]),
        ("4 · Pilot design and investment decision", [
            ("bullet", "Start with one opt-in, low-sensitivity team and existing approved documents."),
            ("bullet", "Baseline active search time, interruptions, response acceptance and time to accepted resolution."),
            ("bullet", "Track eligibility, adoption, supported-answer rate, citation correctness, abstention, correction, escalation, latency, cost per accepted resolution and access violations."),
            ("bullet", "Keep a human fallback, review source ACLs, and report zero access violations as a hard gate."),
            ("p", "Decision gate: proceed only if active time savings are observed without unacceptable access, trust or correction issues; compare with a Copilot Studio workflow before funding a standalone product."),
            ("p", "Do not claim company-wide financial returns from this estimate. Replace assumptions with measured pilot results and include actual Microsoft licensing, security operations and support costs."),
        ]),
    ]
    simple_pdf("HumanAPI_Business_Case.pdf", "Business case: time value and costs", pages)


def build_competitive():
    pages = [
        ("1 · Executive view", [
            ("p", "Many products already find company files, summarize them, show sources and respect file permissions. Human API should not be presented as a new search engine. Microsoft's Autopilot and SharePoint knowledge announcements also overlap with the broad idea of helping work continue when a colleague is away."),
            ("p", "Example: Alice asks why a project setting changed. Bob has allowed questions about that topic for a limited time. The app shows a source Alice can open and explains what the source does not answer; if the current value is unclear, it sends the question to Carla. Whether teams need this extra workflow is still unproven."),
            ("callout", "Recommendation: try the workflow inside Teams or Copilot Studio, reuse Microsoft 365 search and file permissions, and do not build another general enterprise search product."),
        ]),
        ("2 · Capability matrix", [
            ("table", [["What it can do", "Human API demo", "Microsoft tools", "Glean / Moveworks"],
                       ["Find answers and show sources", "Word matching on fictional files", "Copilot and SharePoint knowledge features", "Search and cited answers described"],
                       ["Check who can open each file", "Checks requester before search", "Uses Microsoft file permissions", "Uses connected-app permissions"],
                       ["Run a background agent", "No", "Autopilot announced; preview status varies", "Agents and workflows available"],
                       ["Owner approval and response timer", "Demonstrated locally", "Can be configured; exact same flow not established", "Exact same flow not established"],
                       ["Bring in a human", "Named owner or backup", "Human assistance in Copilot Studio", "Smart handoff / workflow features"],
                       ["Approve a correction before reuse", "Demonstrated locally", "Could be modeled; exact same flow not established", "Exact same flow not established"],
                       ["Public comparable pricing", "None", "License/capacity/usage terms", "Quote required"]]),
        ]),
        ("3 · Evidence by platform", [
            ("h", "Microsoft"),
            ("p", "Microsoft's 25 Sep 2026 announcement describes Autopilot (formerly Scout) as a persistent, proactive personal agent; the announcement says private preview is expanding at month end. Microsoft announced a Copilot-curated team knowledge base in SharePoint Frontier Preview on 14 Oct, and began Copilot in SharePoint's general-availability rollout on 30 Sep. SharePoint agents use source permissions. Copilot Studio supports a human assistance node. These products overlap with work that continues in the background, but the reviewed public pages do not establish the exact owner-approved question and response-timer flow shown in this demo."),
            ("h", "Glean"),
            ("p", "Glean describes cross-app search constrained by existing permissions, grounded answers with citations, people discovery, agent workflows and use in work surfaces. Its public materials overlap heavily with enterprise knowledge search and agent building."),
            ("h", "Moveworks"),
            ("p", "Moveworks documents multi-source search and summarization, source references, conversational Teams/Slack/Google Chat experiences, action plugins and Smart Handoff. Its docs say feedback does not automatically alter assistant behavior and plugin selection cannot be guaranteed by a hard rule."),
        ]),
        ("4 · Licensing and sources", [
            ("bullet", "M365 Copilot requires eligible license plans; agentic features may use consumption-based billing."),
            ("bullet", "Copilot Studio requires authorized user/tenant setup and capacity or usage-based billing according to tenant configuration."),
            ("bullet", "SharePoint agents have applicable access/licensing and tenant governance requirements."),
            ("bullet", "Teams apps require app registration, tenant custom app upload/sideloading and public HTTPS endpoint; Graph delegated permissions need consent."),
            ("bullet", "Comparable Glean/Moveworks pricing was not found in reviewed public technical material; obtain vendor quotes."),
            ("p", "Primary sources accessed 26 Oct 2026: Microsoft Autopilot announcement (blogs.microsoft.com/blog/2026/09/25/introducing-the-new-copilot-with-home-code-and-autopilot); SharePoint team knowledge base announcement (techcommunity.microsoft.com/blog/spblog/introducing-a-team-knowledge-base-curated-by-copilot-and-agents/4564203); M365 agents FAQ and SharePoint agents access (learn.microsoft.com); Copilot Studio licensing and human assistance node (learn.microsoft.com); Glean What is Glean and access docs (docs.glean.com); Moveworks Assistant capabilities (help.moveworks.com/ai-assistant/capabilities)."),
            ("p", "Public sources reflect vendor descriptions and preview states at access time. Verify licenses, tenant availability and product behavior in a procurement proof of concept."),
        ]),
    ]
    simple_pdf("HumanAPI_Competitive_Analysis.pdf", "Competitive landscape and product position", pages)


def build_evaluation():
    report = json.loads((ROOT / "evaluations/latest_results.json").read_text())
    pages = [
        ("1 · Scope and method", [
            ("p", "This report records checks on the local demo, using fictional people and files. No real employee data or Microsoft tenant was used. The demo uses word matching and fixed answer templates, not a language model, so these results say nothing about the quality or cost of a language model."),
            ("p", "The backend tests check the question workflow and access rules. A separate exercise uses 36 written example questions to check whether the search finds the expected accessible file or returns no file when evidence is absent. It is a small word-matching check—not a test of whether an answer is truly correct, whether search understands paraphrases, or whether a production system is secure."),
            ("p", "Run the checks with `make test` and `make eval`. Latest local run: 17 backend tests passed; all 36 search examples passed. The web interface also passed its type check and production build."),
            ("callout", "Example: the question about the 120-second timeout should find ADR-024. The word-matching check confirms that the file was found; it does not prove that the source fully answers whether 120 seconds is still needed."),
        ]),
        ("2 · Observed results", [
            ("table", [["Check", "Result", "What it means"],
                       ["Backend workflow tests", "17 passed", "The tested local rules behaved as expected"],
                       ["Written search examples", f"{report['passed']}/{report['cases']} passed", "32 expected files found; 4 no-evidence questions returned none"],
                       ["File-access violations", "0 in these tests", "Does not prove a real system is secure"],
                       ["Response time, cost, accepted answers", "Not measured", "No live service or user pilot"]]),
            ("h", "Behavior covered by tests"),
            ("bullet", "Supported partial answer includes ADR-024 citation and limitation; restricted budget request denies without title/value."),
            ("bullet", "Low evidence escalates to backup; conflicting timeout documents require clarification."),
            ("bullet", "Human answer before deadline cancels pending job; retry limits stop repeated failure."),
            ("bullet", "Correction remains pending until owner approves; revocation suppresses previously returned answer and citations."),
            ("bullet", "Prompt injection fixture is treated as untrusted text; Restricted labels require explicit ACL."),
        ]),
        ("3 · Interpretation and next evaluation", [
            ("p", "The 100% authored retrieval result is expected for a small synthetic vocabulary and should not be read as general accuracy. Lexical overlap misses paraphrases and can retrieve irrelevant material. Four denial cases do not establish robust resistance to inference or side channels. Prompt-injection sanitation uses a short marker list and is not a complete defense."),
            ("h", "Before a pilot, measure"),
            ("bullet", "Expert-labeled evidence recall and citation precision on representative, non-sensitive questions."),
            ("bullet", "Claim entailment, unsupported-answer rate, abstention precision and contradiction handling with blinded reviewers."),
            ("bullet", "Access regression across users, groups, sensitivity labels, revocation, deletion, tenant boundaries and cached responses."),
            ("bullet", "Concurrency, restart/recovery, notification idempotency, Graph throttling and unavailable service behavior."),
            ("bullet", "Active work time before/after, accepted resolution time, interruption reduction, escalation/correction rates and cost per accepted result."),
            ("callout", "Pilot safety gate: zero unauthorized disclosure. Stop and remediate on any access violation; do not trade that gate against answer coverage."),
        ]),
    ]
    simple_pdf("HumanAPI_Evaluation_Report.pdf", "Evaluation report", pages)


if __name__ == "__main__":
    build_docx()
    build_idea_pdf()
    build_business_case()
    build_competitive()
    build_evaluation()
    copies = {
        "docs/TECHNICAL_DESIGN.md": ("HumanAPI_Technical_Design.md", {
            "(READER_GUIDE.md)": "(HumanAPI_Reader_Guide.md)",
            "(DEPENDENCIES.md)": "(HumanAPI_Dependencies.md)",
        }),
        "docs/SECURITY_ARCHITECTURE.md": ("HumanAPI_Security_Architecture.md", {
            "(READER_GUIDE.md)": "(HumanAPI_Reader_Guide.md)",
        }),
        "docs/MICROSOFT_GRAPH_PERMISSIONS.md": ("HumanAPI_Microsoft_Graph_Permissions.md", {
            "(READER_GUIDE.md)": "(HumanAPI_Reader_Guide.md)",
        }),
        "docs/DEPENDENCIES.md": ("HumanAPI_Dependencies.md", {}),
        "docs/READER_GUIDE.md": ("HumanAPI_Reader_Guide.md", {
            "../deliverables/HumanAPI_Idea_Document.pdf": "HumanAPI_Idea_Document.pdf",
            "../README.md": "README.md",
            "../DEMO_GUIDE.md": "DEMO_GUIDE.md",
            "../deliverables/HumanAPI_Business_Case.pdf": "HumanAPI_Business_Case.pdf",
            "(TECHNICAL_DESIGN.md)": "(HumanAPI_Technical_Design.md)",
            "(SECURITY_ARCHITECTURE.md)": "(HumanAPI_Security_Architecture.md)",
            "(MICROSOFT_GRAPH_PERMISSIONS.md)": "(HumanAPI_Microsoft_Graph_Permissions.md)",
            "../deliverables/HumanAPI_Evaluation_Report.pdf": "HumanAPI_Evaluation_Report.pdf",
            "../PROGRESS.md": "README.md#what-is-implemented",
            "(COMPETITIVE_ANALYSIS.md)": "(HumanAPI_Competitive_Analysis.pdf)",
            "Start with this guide and `README.md`; use `AGENTS.md` for coding boundaries.":
                "Start with this guide. The source repository's `AGENTS.md` contains coding boundaries.",
        }),
        "DEMO_GUIDE.md": ("DEMO_GUIDE.md", {
            "[reader guide](docs/READER_GUIDE.md)": "[reader guide](HumanAPI_Reader_Guide.md)",
            "[AGENTS.md](AGENTS.md)": "the source repository's AGENTS.md",
            "the [reader guide](docs/READER_GUIDE.md)": "the [reader guide](HumanAPI_Reader_Guide.md)",
        }),
    }
    for source_name, (target_name, replacements) in copies.items():
        content = (ROOT / source_name).read_text()
        for old, new in replacements.items():
            content = content.replace(old, new)
        if target_name == "DEMO_GUIDE.md":
            content = content.replace(
                "This demo uses fictional people and project files.",
                "This walkthrough is for the source-repository demo; this document bundle does not include runnable app code. "
                "It uses fictional people and project files.",
            )
            content = content.replace(
                "Follow the setup in [README.md](README.md) and run `make dev`.",
                "In the full source repository, follow the setup in its root README.md and run `make dev`.",
            )
        (OUT / target_name).write_text(content)

    (OUT / "README.md").write_text("""# Human API deliverables

Human API is a concept for handling a work question when the colleague who knows is unavailable. The owner chooses a limited policy; the requester must still be allowed to open each source; the system shows evidence and limits or asks a person to help.

**Example:** Alice asks why a service timeout changed from 90 to 120 seconds. Bob opted in for Project Alpha. A written decision explains why the change was approved, but not whether it is still needed. Human API shows that partial answer with its source; it does not speak as Bob.

This bundle explains the concept and includes a local demo walkthrough, but it does not contain runnable application code. The demo uses fictional people and files, word matching and fixed templates. It uses no language model and is not connected to a Microsoft tenant. There is no Teams bot.

## Read or present

- [Reader guide](HumanAPI_Reader_Guide.md) — plain-language overview, example, glossary and future AI-agent guardrails.
- [Pitch deck (PowerPoint)](HumanAPI_Pitch_Deck.pptx) — editable slides; speaker notes contain technical detail and source links.
- [Pitch deck (PDF)](HumanAPI_Pitch_Deck.pdf) — matching 15-slide handout; see the companion PowerPoint for notes and source links.
- [Business review (PowerPoint)](HumanAPI_Business_Review.pptx) and [PDF](HumanAPI_Business_Review.pdf) — current concept/MVP status, synthetic evaluation, illustrative value, decisions and a proposed pilot path.
- [Idea document (PDF)](HumanAPI_Idea_Document.pdf) and [editable DOCX](HumanAPI_Idea_Document.docx).
- [Five-minute demo walkthrough](DEMO_GUIDE.md) — requires the full source repository to run.

## Supporting documents

- [Business case](HumanAPI_Business_Case.pdf) — illustrative time-value assumptions, clearly marked as unmeasured.
- [Competitive analysis](HumanAPI_Competitive_Analysis.pdf).
- [Evaluation report](HumanAPI_Evaluation_Report.pdf) — local test scope and limitations.
- [Technical design](HumanAPI_Technical_Design.md).
- [Security architecture](HumanAPI_Security_Architecture.md).
- [Microsoft Graph permissions](HumanAPI_Microsoft_Graph_Permissions.md) — integration plan; not tenant-tested.
- [Dependencies and licenses](HumanAPI_Dependencies.md).

The code and additional project notes are in the full Human API source repository.
""")
    print("Generated DOCX and 4 PDF reports in deliverables/")
